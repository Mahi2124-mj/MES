# ════════════════════════════════════════════════════════════════
# routers/py_stations.py
# ════════════════════════════════════════════════════════════════
"""
Poka-Yoke STATIONS  →  /api/py-stations/...

Machine-station model (2026-07, per Vivek's spec):
    Zone → production Line → its machines (from mes_machines) → assign a
    STATION CODE (SS_01, SS_02 … user's choice) to each machine.
    Machines that share a code across lines = one "station group" (paired).
    PYs are defined per station group → the same PYs apply to every paired
    line's machine.  Add / delete PY per station.

Tables (additive, no collector touch):
    mes_machine_station  — machine_id (→ mes_machines) → station_code
    mes_station_py       — (zone_name, station_code) → PY definitions
Machines come from mes_machines (per-line directory: machine_no = LINE_SS_NN).
Production lines come from mes_lines (Production Panel).  A production line
maps to its mes_machines rows by normalised line_name (YNC-SS ↔ YNC_SS).
"""

from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from fastapi.responses import Response
from pydantic import BaseModel
from typing import Optional

from database import get_conn, dict_cursor
from auth import get_current_user, require_admin

router = APIRouter(prefix="/api/py-stations", tags=["py-stations"])

_NORM = "upper(replace(replace(%s,'-','_'),' ','_'))"


def _norm(s: str) -> str:
    return (s or "").upper().replace("-", "_").replace(" ", "_")


def _zone_name(cur, zone_id: int) -> str:
    cur.execute("SELECT zone_name FROM mes_zones WHERE id=%s", (zone_id,))
    r = cur.fetchone()
    return r["zone_name"] if r else str(zone_id)


# ══════════════════════════════════════════════════════════════
# READ — navigation
# ══════════════════════════════════════════════════════════════
@router.get("/zones")
def list_zones(user=Depends(get_current_user)):
    """Production zones that have production lines."""
    with get_conn() as conn:
        cur = dict_cursor(conn)
        cur.execute("""
            SELECT z.id                    AS "zoneId",
                   z.zone_name             AS "zoneName",
                   count(DISTINCT l.id)    AS "lineCount"
            FROM mes_zones z
            JOIN mes_lines l ON l.zone_id = z.id
            GROUP BY z.id, z.zone_name
            ORDER BY z.zone_name
        """)
        return cur.fetchall()


@router.get("/lines")
def list_lines(zone_id: int = Query(...), user=Depends(get_current_user)):
    with get_conn() as conn:
        cur = dict_cursor(conn)
        cur.execute("""
            SELECT id AS "lineId", line_name AS "lineName", line_code AS "lineCode"
            FROM mes_lines WHERE zone_id=%s ORDER BY line_name
        """, (zone_id,))
        return cur.fetchall()


@router.get("/line/{line_id}/machines")
def line_machines(line_id: int, user=Depends(get_current_user)):
    """Machines (mes_machines) that belong to a production line, with the
    station code assigned to each (if any)."""
    with get_conn() as conn:
        cur = dict_cursor(conn)
        cur.execute("SELECT line_name FROM mes_lines WHERE id=%s", (line_id,))
        r = cur.fetchone()
        if not r:
            raise HTTPException(404, "line not found")
        cur.execute(f"""
            SELECT m.id,
                   m.machine_no     AS "machineNo",
                   m.machine_name   AS "machineName",
                   m.serial_no      AS "serialNo",
                   s.station_code   AS "stationCode"
            FROM mes_machines m
            LEFT JOIN mes_machine_station s ON s.machine_id = m.id
            WHERE {_NORM % 'm.line_name'} = {_NORM % '%s'}
              AND COALESCE(m.is_active, true)
            ORDER BY m.serial_no NULLS LAST, m.machine_no
        """, (r["line_name"],))
        return cur.fetchall()


@router.get("/groups")
def station_groups(zone_id: int = Query(...), user=Depends(get_current_user)):
    """Station groups (by code) across a zone's production lines — the
    pairing view.  Each code → the machines (across lines) that share it."""
    with get_conn() as conn:
        cur = dict_cursor(conn)
        cur.execute("SELECT id, line_name FROM mes_lines WHERE zone_id=%s", (zone_id,))
        norms = { _norm(l["line_name"]) for l in cur.fetchall() }
        cur.execute("""
            SELECT s.station_code, m.id AS machine_id, m.line_name, m.machine_no, m.machine_name
            FROM mes_machine_station s
            JOIN mes_machines m ON m.id = s.machine_id
            ORDER BY s.station_code, m.line_name
        """)
        groups = {}
        for r in cur.fetchall():
            if _norm(r["line_name"]) not in norms:
                continue
            g = groups.setdefault(r["station_code"], {"stationCode": r["station_code"], "machines": []})
            g["machines"].append({
                "machineId":   r["machine_id"],
                "lineName":    r["line_name"],
                "machineNo":   r["machine_no"],
                "machineName": r["machine_name"],
            })
        return sorted(groups.values(), key=lambda x: x["stationCode"])


# ══════════════════════════════════════════════════════════════
# WRITE — assign station code
# ══════════════════════════════════════════════════════════════
class CodeBody(BaseModel):
    station_code: Optional[str] = None


@router.put("/machine/{machine_id}/code")
def set_code(machine_id: int, body: CodeBody, admin=Depends(require_admin)):
    """Assign / clear the station code of a machine (blank = clear)."""
    code = (body.station_code or "").strip().upper()
    with get_conn() as conn:
        cur = dict_cursor(conn)
        cur.execute("SELECT id FROM mes_machines WHERE id=%s", (machine_id,))
        if not cur.fetchone():
            raise HTTPException(404, "machine not found")
        if not code:
            cur.execute("DELETE FROM mes_machine_station WHERE machine_id=%s", (machine_id,))
            return {"ok": True, "code": None}
        cur.execute("""
            INSERT INTO mes_machine_station(machine_id, station_code, updated_at)
            VALUES (%s,%s,now())
            ON CONFLICT (machine_id) DO UPDATE
              SET station_code=EXCLUDED.station_code, updated_at=now()
        """, (machine_id, code))
        return {"ok": True, "code": code}


# ══════════════════════════════════════════════════════════════
# STATION PYs  (shared across every paired line's machine)
# ══════════════════════════════════════════════════════════════
@router.get("/pys")
def list_pys(zone_id: int = Query(...), station_code: str = Query(...),
             user=Depends(get_current_user)):
    with get_conn() as conn:
        cur = dict_cursor(conn)
        zn = _zone_name(cur, zone_id)
        cur.execute("""
            SELECT id,
                   py_name         AS "pyName",
                   d_bit           AS "dBit",
                   sensing_bits    AS "sensingBits",
                   register_count  AS "registerCount",
                   desired_value   AS "desiredValue",
                   desired_value_2 AS "desiredValue2"
            FROM mes_station_py
            WHERE zone_name=%s AND station_code=%s AND is_active
            ORDER BY id
        """, (zn, station_code.strip().upper()))
        return cur.fetchall()


class StationPyCreate(BaseModel):
    zone_id:         int
    station_code:    str
    py_name:         str
    d_bit:           Optional[str] = None
    sensing_bits:    Optional[str] = None
    register_count:  Optional[int] = 1
    desired_value:   Optional[int] = None
    desired_value_2: Optional[int] = None


@router.post("/pys", status_code=201)
def add_py(body: StationPyCreate, admin=Depends(require_admin)):
    name = (body.py_name or "").strip()
    code = (body.station_code or "").strip().upper()
    if not name or not code:
        raise HTTPException(400, "py_name and station_code required")
    with get_conn() as conn:
        cur = dict_cursor(conn)
        zn = _zone_name(cur, body.zone_id)
        cur.execute("""
            INSERT INTO mes_station_py
                (zone_name, station_code, py_name, d_bit, sensing_bits,
                 register_count, desired_value, desired_value_2)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s)
            RETURNING id
        """, (zn, code, name, body.d_bit, body.sensing_bits,
              body.register_count or 1, body.desired_value, body.desired_value_2))
        return {"id": cur.fetchone()["id"], "ok": True}


@router.delete("/pys/{py_id}")
def delete_py(py_id: int, admin=Depends(require_admin)):
    with get_conn() as conn:
        cur = dict_cursor(conn)
        cur.execute("DELETE FROM mes_station_py WHERE id=%s RETURNING id", (py_id,))
        if not cur.fetchone():
            raise HTTPException(404, "py not found")
        return {"ok": True}


# ══════════════════════════════════════════════════════════════
# PER-MACHINE override — a station PY can have a different desired
# output on each paired line's machine, or be turned off there.
# ══════════════════════════════════════════════════════════════
@router.get("/pys/{py_id}/machines")
def py_machine_outputs(py_id: int, user=Depends(get_current_user)):
    """Every paired-line machine for this station PY, with its effective
    output — the per-machine override if set, else the PY's default."""
    with get_conn() as conn:
        cur = dict_cursor(conn)
        cur.execute("""SELECT zone_name, station_code, d_bit, desired_value, desired_value_2
                       FROM mes_station_py WHERE id=%s""", (py_id,))
        py = cur.fetchone()
        if not py:
            raise HTTPException(404, "py not found")
        # normalised production-line names of the PY's zone (to scope the group)
        cur.execute("SELECT id FROM mes_zones WHERE zone_name=%s", (py["zone_name"],))
        zr = cur.fetchone()
        line_norms = set()
        if zr:
            cur.execute("SELECT line_name FROM mes_lines WHERE zone_id=%s", (zr["id"],))
            line_norms = { _norm(l["line_name"]) for l in cur.fetchall() }
        cur.execute("""
            SELECT m.id AS machine_id, m.line_name, m.machine_no, m.machine_name,
                   o.enabled, o.d_bit AS o_dbit, o.desired_value AS o_dv,
                   o.desired_value_2 AS o_dv2, o.id AS ov_id
            FROM mes_machine_station s
            JOIN mes_machines m ON m.id = s.machine_id
            LEFT JOIN mes_station_py_machine o
                   ON o.station_py_id=%s AND o.machine_id=m.id
            WHERE s.station_code=%s
            ORDER BY m.line_name
        """, (py_id, py["station_code"]))
        out = []
        for r in cur.fetchall():
            if line_norms and _norm(r["line_name"]) not in line_norms:
                continue
            out.append({
                "machineId":     r["machine_id"],
                "lineName":      r["line_name"],
                "machineNo":     r["machine_no"],
                "machineName":   r["machine_name"],
                "enabled":       (r["enabled"] if r["enabled"] is not None else True),
                "dBit":          (r["o_dbit"] if r["o_dbit"] is not None else py["d_bit"]),
                "desiredValue":  (r["o_dv"]   if r["o_dv"]   is not None else py["desired_value"]),
                "desiredValue2": (r["o_dv2"]  if r["o_dv2"]  is not None else py["desired_value_2"]),
                "hasOverride":   r["ov_id"] is not None,
                "defaultDbit":   py["d_bit"],
                "defaultDv":     py["desired_value"],
                "defaultDv2":    py["desired_value_2"],
            })
        return out


class MachineOverride(BaseModel):
    enabled:         Optional[bool] = True
    d_bit:           Optional[str] = None
    desired_value:   Optional[int] = None
    desired_value_2: Optional[int] = None


@router.put("/pys/{py_id}/machine/{machine_id}")
def set_machine_override(py_id: int, machine_id: int, body: MachineOverride,
                         admin=Depends(require_admin)):
    """Set this machine's own output for the PY (or turn the PY off here)."""
    with get_conn() as conn:
        cur = dict_cursor(conn)
        cur.execute("SELECT id FROM mes_station_py WHERE id=%s", (py_id,))
        if not cur.fetchone():
            raise HTTPException(404, "py not found")
        cur.execute("""
            INSERT INTO mes_station_py_machine
                (station_py_id, machine_id, enabled, d_bit, desired_value, desired_value_2, updated_at)
            VALUES (%s,%s,%s,%s,%s,%s, now())
            ON CONFLICT (station_py_id, machine_id) DO UPDATE SET
                enabled=EXCLUDED.enabled, d_bit=EXCLUDED.d_bit,
                desired_value=EXCLUDED.desired_value,
                desired_value_2=EXCLUDED.desired_value_2, updated_at=now()
        """, (py_id, machine_id,
              body.enabled if body.enabled is not None else True,
              (body.d_bit or "").strip() or None,
              body.desired_value, body.desired_value_2))
        return {"ok": True}


@router.delete("/pys/{py_id}/machine/{machine_id}")
def clear_machine_override(py_id: int, machine_id: int, admin=Depends(require_admin)):
    """Reset a machine back to the PY's default output."""
    with get_conn() as conn:
        cur = dict_cursor(conn)
        cur.execute("DELETE FROM mes_station_py_machine WHERE station_py_id=%s AND machine_id=%s",
                    (py_id, machine_id))
        return {"ok": True}


# ══════════════════════════════════════════════════════════════
# EXCEL — template (every machine + its current code) + bulk import
# ══════════════════════════════════════════════════════════════
def _xlsx(wb, fname: str) -> Response:
    import io
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


@router.get("/template", tags=["import-export"])
def download_template(user=Depends(get_current_user)):
    """Excel with EVERY machine (from mes_machines) + its current station
    code — fill the Station Code column and re-import."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    with get_conn() as conn:
        cur = dict_cursor(conn)
        cur.execute("""
            SELECT m.zone_name, m.line_name, m.machine_no, m.machine_name, s.station_code
            FROM mes_machines m
            LEFT JOIN mes_machine_station s ON s.machine_id = m.id
            WHERE COALESCE(m.is_active, true)
            ORDER BY m.zone_name, m.line_name, m.serial_no NULLS LAST, m.machine_no
        """)
        rows = cur.fetchall()
    wb = Workbook(); ws = wb.active; ws.title = "Machines"
    ws.append(["Zone", "Line", "Machine No", "Machine Name", "Station Code"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="16A34A")
    for r in rows:
        ws.append([r["zone_name"], r["line_name"], r["machine_no"],
                   r["machine_name"], r["station_code"] or ""])
    ws2 = wb.create_sheet("Instructions")
    ws2.append(["Column", "Notes"])
    ws2[1][0].font = ws2[1][1].font = Font(bold=True, color="FFFFFF")
    for c in ws2[1]:
        c.fill = PatternFill("solid", fgColor="0F172A")
    ws2.append(["Machine No", "Machine ka unique code (YNC_SS_01) — change MAT karo, isi se match hota hai"])
    ws2.append(["Station Code", "SS_01, SS_02… bharo. Same code across lines = paired station. Blank = code hatao."])
    ws2.append(["Zone / Line / Machine Name", "Sirf reference — import inhe ignore karta hai (Machine No + Station Code use hote)"])
    for w in (ws, ws2):
        for col in w.columns:
            ml = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            w.column_dimensions[col[0].column_letter].width = min(ml + 3, 60)
    return _xlsx(wb, "py_stations_machines_template.xlsx")


@router.post("/import", tags=["import-export"])
async def import_codes(file: UploadFile = File(...), admin=Depends(require_admin)):
    """Bulk-assign station codes from the filled template.  Matches each row
    by Machine No; blank Station Code clears that machine's code."""
    import io
    from openpyxl import load_workbook
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(400, "Upload an .xlsx file")
    contents = await file.read()
    try:
        wb = load_workbook(io.BytesIO(contents), data_only=True)
    except Exception:
        raise HTTPException(400, "Could not read the Excel file")
    ws = wb["Machines"] if "Machines" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"updated": 0, "cleared": 0, "errors": ["empty sheet"]}
    header = [(str(h).strip().lower() if h is not None else "") for h in rows[0]]
    try:
        ci_no   = header.index("machine no")
        ci_code = header.index("station code")
    except ValueError:
        raise HTTPException(400, "Sheet mein 'Machine No' aur 'Station Code' columns chahiye")
    updated = cleared = 0
    errors = []
    with get_conn() as conn:
        cur = dict_cursor(conn)
        for i, row in enumerate(rows[1:], start=2):
            if not row or ci_no >= len(row) or row[ci_no] is None:
                continue
            mno  = str(row[ci_no]).strip()
            code = (str(row[ci_code]).strip().upper()
                    if ci_code < len(row) and row[ci_code] is not None else "")
            if not mno:
                continue
            cur.execute("SELECT id FROM mes_machines WHERE machine_no=%s", (mno,))
            mr = cur.fetchone()
            if not mr:
                errors.append(f"row {i}: machine '{mno}' not found")
                continue
            if code:
                cur.execute("""
                    INSERT INTO mes_machine_station(machine_id, station_code, updated_at)
                    VALUES (%s,%s,now())
                    ON CONFLICT (machine_id) DO UPDATE
                      SET station_code=EXCLUDED.station_code, updated_at=now()
                """, (mr["id"], code))
                updated += 1
            else:
                cur.execute("DELETE FROM mes_machine_station WHERE machine_id=%s", (mr["id"],))
                cleared += 1
    return {"updated": updated, "cleared": cleared, "errors": errors[:25]}
