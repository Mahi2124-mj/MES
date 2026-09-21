"""
routers/machine_master.py
=========================
Machine Master (the user's Excel "Machines Master") → DB table
mes_machine_master + LIVE status.  Separate from routers/machines.py
(which is the zones.json lookup + process graphs) to avoid any clash.

Source = the user's xlsx, imported into Postgres (replaced on each upload).
A background poller ICMP-pings every machine IP AND camera IP so the
Admin → Machine Master tab shows green/red per machine + per camera.

Endpoints:
  GET  /api/machine-master          → rows + ip_up / cam_up (live)
  POST /api/machine-master/upload   → admin uploads the xlsx → replaces table
"""
import io
import time
import asyncio
import threading
import subprocess

from fastapi import APIRouter, Depends, UploadFile, File, HTTPException

from auth import get_current_user, require_admin
from database import get_conn, dict_cursor
from ddl_once import once

router = APIRouter(prefix="/api/machine-master", tags=["machine-master"])

_POLL_EVERY = 60
_status = {}                       # ip -> bool (reachable)
_status_lock = threading.Lock()


@once
def _ensure_table():
    with get_conn() as conn:
        conn.cursor().execute("""
            CREATE TABLE IF NOT EXISTS mes_machine_master (
              id            SERIAL PRIMARY KEY,
              sort_order    INT,
              zone          TEXT,
              line          TEXT,
              machine_no    TEXT,
              machine_name  TEXT,
              ip            TEXT,
              port          TEXT,
              camera_ip     TEXT,
              data_register TEXT
            )""")


def _rows():
    with get_conn() as conn:
        cur = dict_cursor(conn)
        cur.execute("SELECT * FROM mes_machine_master ORDER BY sort_order, id")
        return [dict(r) for r in cur.fetchall()]


def _import_xlsx(data) -> int:
    """Parse the 'Machines Master' sheet (row1 = title, row2 = headers, row3+ =
    data) and REPLACE the table.  `data` = file path or raw bytes.
    Status / Machine Active columns are ignored."""
    import openpyxl
    if isinstance(data, (bytes, bytearray)):
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    else:
        wb = openpyxl.load_workbook(data, read_only=True, data_only=True)
    ws = wb["Machines Master"] if "Machines Master" in wb.sheetnames else wb.worksheets[0]

    header, body = None, []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 1:
            header = [str(c).strip().lower() if c is not None else "" for c in r]
        elif i >= 2 and any(c is not None for c in r):
            body.append(r)
    wb.close()
    if not header:
        raise ValueError("header row (row 2) not found")

    def gi(*names):
        for n in names:
            if n in header:
                return header.index(n)
        return None
    idx = {
        "zone": gi("zone"), "line": gi("line"),
        "machine_no": gi("machine no", "machine_no"),
        "machine_name": gi("machine name", "machine_name"),
        "ip": gi("ip"), "port": gi("port"),
        "camera_ip": gi("camera ip", "camera_ip"),
        "data_register": gi("data register", "data_register"),
    }

    out = []
    for so, r in enumerate(body):
        def val(k):
            i = idx[k]
            return "" if (i is None or i >= len(r) or r[i] is None) else str(r[i]).strip()
        if not (val("ip") or val("machine_no") or val("machine_name")):
            continue
        out.append((so, val("zone"), val("line"), val("machine_no"), val("machine_name"),
                    val("ip"), val("port"), val("camera_ip"), val("data_register")))

    _ensure_table()
    with get_conn() as conn:
        c = conn.cursor()
        c.execute("TRUNCATE mes_machine_master RESTART IDENTITY")
        c.executemany("""INSERT INTO mes_machine_master
            (sort_order, zone, line, machine_no, machine_name, ip, port, camera_ip, data_register)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)""", out)
    return len(out)


def _ping(ip: str) -> bool:
    try:
        r = subprocess.run(["ping", "-n", "1", "-w", "700", ip],
                           capture_output=True, text=True, timeout=4)
        return r.returncode == 0 and "ttl=" in (r.stdout or "").lower()
    except Exception:
        return False


async def _poll():
    ips = set()
    for r in _rows():
        for k in ("ip", "camera_ip"):
            v = (r.get(k) or "").strip()
            if v:
                ips.add(v)
    if not ips:
        return
    sem = asyncio.Semaphore(40)
    res = {}
    async def one(ip):
        async with sem:
            res[ip] = await asyncio.to_thread(_ping, ip)
    await asyncio.gather(*[one(ip) for ip in ips])
    with _status_lock:
        _status.clear()
        _status.update(res)


def _poller():
    loop = asyncio.new_event_loop()
    asyncio.set_event_loop(loop)
    print(f"[MMASTER] machine-master poller started (every {_POLL_EVERY}s)")
    while True:
        try:
            loop.run_until_complete(_poll())
        except Exception as exc:
            print(f"[MMASTER] error: {exc}", flush=True)
        time.sleep(_POLL_EVERY)


def start_machine_master_poller():
    try:
        _ensure_table()
    except Exception as exc:
        print(f"[MMASTER] table init failed: {exc}", flush=True)
    threading.Thread(target=_poller, name="mmaster-poll", daemon=True).start()


# ── endpoints ────────────────────────────────────────────────────────
@router.get("")
def list_machine_master(user=Depends(get_current_user)):
    rows = _rows()
    with _status_lock:
        st = dict(_status)
    for r in rows:
        ip = (r.get("ip") or "").strip()
        cam = (r.get("camera_ip") or "").strip()
        r["ip_up"] = st.get(ip) if ip in st else None
        r["cam_up"] = st.get(cam) if cam in st else None
    return {"machines": rows, "count": len(rows)}


@router.post("/upload")
async def upload_machine_master(file: UploadFile = File(...), admin=Depends(require_admin)):
    data = await file.read()
    try:
        n = _import_xlsx(data)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Import failed: {exc}")
    return {"ok": True, "count": n}
