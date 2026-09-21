"""routers/video_coverage.py — Video Coverage log + the read-only Video Agent.

Per cycle: was a video clip made, and if not, why — line-wise and machine-wise,
with the full list of cycles that have no clip.  A read-only agent watches it.

Why the reason has to be written down WHEN it happens: the CMS deletes every
camera's .ts at each shift start and the clip archiver gives up on a cycle after
42 minutes, so an hour later nothing on disk can say whether the camera was
dead, hung, or simply wiped.  So a small tracker samples the cameras every
minute and keeps their state as intervals, and each cycle is judged once, 45
minutes after it ended, against that record.

Everything here is READ-ONLY towards the plant:
  * cameras  — ICMP ping only (for cameras that are not recording).  Never a
               TCP/RTSP connection: most cameras accept ONE session, and a
               probe can knock a live recorder off.
  * CMS      — a TCP connect to its own port on localhost, and a directory
               listing of the video folder.  Nothing is started or stopped.
  * clips    — existence checks in the clip archive.
The Video Agent only raises findings and Inbox notices.  It has NO code path to
restart, kill, reconfigure or write to anything outside its own tables.

Tracking starts at go-live (operator chose no backfill): cycles before the first
evaluation are not judged.
"""

import io
import json
import os
import re
import socket
import subprocess
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse

from auth import get_current_user
from database import get_conn, dict_cursor
from ddl_once import once

router = APIRouter(prefix="/api/video-coverage", tags=["video-coverage"])

MATURE_MIN   = int(os.environ.get("VCOV_MATURE_MIN", "45"))   # archiver window is 42 min
SAMPLE_S     = 60
EVAL_S       = 120
AGENT_S      = 300
RECHECK_H    = 3          # a late clip (archiver backlog) still turns a "missing" into a clip
RETAIN_DAYS  = 30         # same as the clip archive
SLACK_S      = 75         # a sample row stands for the minute around it
NOTIFY_MIN_GAP_S = 15 * 60

_HERE = os.path.dirname(os.path.abspath(__file__))
_CMS_DIRS = (
    os.environ.get("CMS_BACKEND_DIR", ""),
    os.path.join(_HERE, "..", "..", "..", "..", "New folder (2)", "New folder (2)", "backend"),
    "/run/media/server/5b4e6f01-4c04-4bc3-b34e-b0df376c067f/server backup/"
    "D DRIVE/EOL/EOL/New folder (2)/New folder (2)/backend",
)
_TABLE_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
_TS_RE = re.compile(r"^cam_(.+)_(\d{13})\.ts$")

REASONS = {
    "no_camera":      "No camera configured",
    "cms_down":       "CMS down",
    "network_down":   "Network / switch down",
    "camera_offline": "Camera offline (no ping)",
    "camera_hung":    "Camera hung (ping OK, no video)",
    "shift_wipe":     "Footage deleted at shift change",
    "clip_failed":    "Clip not cut (camera was recording)",
    "unknown":        "Not tracked (tracker was not running)",
}
# worst first — a cycle overlapping several states takes the worst one
_STATE_RANK = {"cms_down": 0, "camera_offline": 2, "camera_hung": 3, "recording": 9}


# ── schema ────────────────────────────────────────────────────────────────
@once
def _ensure_tables():
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute("""
            CREATE TABLE IF NOT EXISTS mes_vcov_cam_state (
                id          BIGSERIAL PRIMARY KEY,
                camera_id   TEXT NOT NULL,
                ip          TEXT,
                state       TEXT NOT NULL,
                subnet_down BOOLEAN NOT NULL DEFAULT FALSE,
                from_ts     TIMESTAMPTZ NOT NULL,
                to_ts       TIMESTAMPTZ NOT NULL)""")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_vcov_cam_state_cam "
                    "ON mes_vcov_cam_state (camera_id, to_ts)")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_vcov_cam_state_to "
                    "ON mes_vcov_cam_state (to_ts)")
        cur.execute("""
            CREATE TABLE IF NOT EXISTS mes_vcov_hourly (
                record_date  DATE NOT NULL,
                shift_name   TEXT NOT NULL,
                line_id      INT  NOT NULL,
                machine_key  TEXT NOT NULL,
                machine_name TEXT,
                hour_ts      TIMESTAMPTZ NOT NULL,
                cycles       INT NOT NULL DEFAULT 0,
                clips        INT NOT NULL DEFAULT 0,
                missing      INT NOT NULL DEFAULT 0,
                PRIMARY KEY (record_date, shift_name, line_id, machine_key, hour_ts))""")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_vcov_hourly_hour "
                    "ON mes_vcov_hourly (hour_ts)")
        cur.execute("""
            CREATE TABLE IF NOT EXISTS mes_vcov_missing (
                id           BIGSERIAL PRIMARY KEY,
                record_date  DATE NOT NULL,
                shift_name   TEXT NOT NULL,
                line_id      INT  NOT NULL,
                machine_key  TEXT NOT NULL,
                machine_name TEXT,
                cycle_seq    INT  NOT NULL,
                is_ng        BOOLEAN NOT NULL DEFAULT FALSE,
                ts_start     TIMESTAMPTZ,
                ts_end       TIMESTAMPTZ NOT NULL,
                part_code    TEXT,
                ct           NUMERIC,
                camera_id    TEXT,
                camera_ip    TEXT,
                reason       TEXT NOT NULL,
                detail       TEXT,
                evaluated_at TIMESTAMPTZ NOT NULL DEFAULT now(),
                UNIQUE (line_id, machine_key, record_date, shift_name, cycle_seq, is_ng))""")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_vcov_missing_end "
                    "ON mes_vcov_missing (ts_end)")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_vcov_missing_line "
                    "ON mes_vcov_missing (line_id, ts_end)")
        cur.execute("""
            CREATE TABLE IF NOT EXISTS mes_vcov_cursor (
                source  TEXT PRIMARY KEY,
                last_ts TIMESTAMPTZ NOT NULL)""")
        cur.execute("""
            CREATE TABLE IF NOT EXISTS mes_vcov_findings (
                id        BIGSERIAL PRIMARY KEY,
                kind      TEXT NOT NULL,
                fkey      TEXT NOT NULL,
                line_id   INT,
                camera_id TEXT,
                message   TEXT NOT NULL,
                opened_at TIMESTAMPTZ NOT NULL DEFAULT now(),
                last_seen TIMESTAMPTZ NOT NULL DEFAULT now(),
                closed_at TIMESTAMPTZ,
                notified  BOOLEAN NOT NULL DEFAULT FALSE)""")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_vcov_findings_open "
                    "ON mes_vcov_findings (fkey) WHERE closed_at IS NULL")
        cur.execute("""
            CREATE TABLE IF NOT EXISTS mes_vcov_agent (
                id           INT PRIMARY KEY,
                started_at   TIMESTAMPTZ,
                tracking_since TIMESTAMPTZ,
                last_sample  TIMESTAMPTZ,
                last_eval    TIMESTAMPTZ,
                last_agent   TIMESTAMPTZ,
                last_notify  TIMESTAMPTZ,
                last_error   TEXT)""")
        cur.execute("INSERT INTO mes_vcov_agent (id) VALUES (1) ON CONFLICT (id) DO NOTHING")


def _agent_mark(**kw):
    if not kw:
        return
    cols = ", ".join(f"{k} = %s" for k in kw)
    try:
        with get_conn() as conn:
            conn.cursor().execute(f"UPDATE mes_vcov_agent SET {cols} WHERE id = 1",
                                  list(kw.values()))
    except Exception as exc:
        print(f"[VCOV] agent status write failed: {exc}", flush=True)


# ── plant facts (read-only) ─────────────────────────────────────────────────
def _cms_file(name):
    for d in _CMS_DIRS:
        if d and os.path.exists(os.path.join(d, name)):
            return os.path.join(d, name)
    return None


def _load_cameras():
    """{camera_id: ip} from the CMS's cameras.json — ids and IPs only, never
    the credentials that file also holds."""
    p = _cms_file("cameras.json")
    if not p:
        return {}
    try:
        data = json.load(open(p, encoding="utf-8"))
        cams = data.get("cameras", data) if isinstance(data, dict) else data
        return {str(c.get("id")).strip(): (c.get("ip") or "").strip()
                for c in (cams or []) if c.get("id")}
    except Exception:
        return {}


def _shift_boundaries():
    """Start times (HH:MM) of every CMS shift window — the moments the CMS
    deletes camera footage (its shift-wipe)."""
    p = _cms_file("shifts.json")
    out = []
    try:
        for s in (json.load(open(p, encoding="utf-8")).get("shifts") or []):
            st = str(s.get("start") or "")
            if re.match(r"^\d{1,2}:\d{2}$", st):
                h, m = st.split(":")
                out.append((int(h), int(m)))
    except Exception:
        pass
    return out or [(8, 30), (17, 20), (18, 30), (3, 20)]


def _video_root():
    try:
        from routers.clip_archive import VIDEO_ROOT
        return VIDEO_ROOT
    except Exception:
        return "/run/media/server/3ad0fece-b7bc-48b1-8f24-d21bb5153735/eol-data/videos"


def _cms_up():
    s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    s.settimeout(2.0)
    try:
        s.connect(("127.0.0.1", int(os.environ.get("CMS_PORT", "5555"))))
        return True
    except Exception:
        return False
    finally:
        s.close()


def _ping(ip):
    if not ip:
        return False
    try:
        r = subprocess.run(["ping", "-c", "1", "-W", "1", ip],
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                           timeout=4)
        return r.returncode == 0
    except Exception:
        return False


def _machines(cur):
    """Every clip source: each line's main (Final-Inspection) camera and each
    sub-machine's camera — resolved exactly as the clip archiver resolves them."""
    cur.execute("""
        SELECT l.id, l.line_name, l.db_table_name, z.zone_name,
               (SELECT NULLIF(TRIM(pc.nf2_camera_id), '')
                  FROM mes_plc_configs pc
                 WHERE pc.line_id = l.id
                   AND ((l.dashboard_plc_id IS NOT NULL AND pc.id = l.dashboard_plc_id)
                     OR (l.dashboard_plc_id IS NULL AND pc.parent_plc_id IS NULL))
                 ORDER BY (pc.parent_plc_id IS NULL) DESC LIMIT 1) AS cam,
               (SELECT NULLIF(TRIM(pc.machine_name), '')
                  FROM mes_plc_configs pc
                 WHERE pc.line_id = l.id
                   AND ((l.dashboard_plc_id IS NOT NULL AND pc.id = l.dashboard_plc_id)
                     OR (l.dashboard_plc_id IS NULL AND pc.parent_plc_id IS NULL))
                 ORDER BY (pc.parent_plc_id IS NULL) DESC LIMIT 1) AS mname
          FROM mes_lines l
          LEFT JOIN mes_zones z ON z.id = l.zone_id
         WHERE COALESCE(l.is_active, TRUE) AND COALESCE(l.db_table_name, '') <> ''""")
    lines = cur.fetchall() or []
    cur.execute("""SELECT id, line_id, NULLIF(TRIM(machine_name), '') AS mname,
                          NULLIF(TRIM(nf2_camera_id), '') AS cam
                     FROM mes_plc_configs""")
    subs = {r["id"]: r for r in (cur.fetchall() or [])}
    return lines, subs


# ── camera tracker ──────────────────────────────────────────────────────────
_last_sizes = {}          # camera_id -> (file name, size) at the previous sample
_last_rows = {}           # camera_id -> (row id, state, subnet_down)
_lock = threading.Lock()


def _sample():
    _ensure_tables()
    cams = _load_cameras()
    if not cams:
        return
    cms_ok = _cms_up()
    newest = {}
    root = _video_root()
    now_wall = time.time()
    try:
        for n in os.listdir(root):
            m = _TS_RE.match(n)
            if not m:
                continue
            try:
                st = os.stat(os.path.join(root, n))
            except FileNotFoundError:
                continue
            cid = m.group(1)
            if cid not in newest or st.st_mtime > newest[cid][2]:
                newest[cid] = (n, st.st_size, st.st_mtime)
    except Exception:
        pass

    states = {}
    if not cms_ok:
        states = {cid: "cms_down" for cid in cams}
    else:
        idle = []
        for cid in cams:
            nw = newest.get(cid)
            prev = _last_sizes.get(cid)
            grew = bool(nw and prev and nw[0] == prev[0] and nw[1] > prev[1])
            fresh = bool(nw and now_wall - nw[2] < 20)
            if grew or (fresh and (prev is None or nw[0] != prev[0])):
                states[cid] = "recording"
            else:
                idle.append(cid)
        with ThreadPoolExecutor(32) as ex:
            for cid, ok in zip(idle, ex.map(lambda c: _ping(cams.get(c)), idle)):
                states[cid] = "camera_hung" if ok else "camera_offline"
    for cid, nw in newest.items():
        _last_sizes[cid] = (nw[0], nw[1])

    # a whole /24 going dark together is the network, not 20 separate cameras
    by_net = {}
    for cid, ip in cams.items():
        net = ".".join(ip.split(".")[:3]) if ip else ""
        a = by_net.setdefault(net, [0, 0])
        a[0] += 1
        a[1] += states.get(cid) in ("camera_offline", "camera_hung")
    dark = {net for net, (n, d) in by_net.items() if net and n >= 3 and d / n >= 0.8}

    with _lock, get_conn() as conn:
        cur = conn.cursor()
        keep = []
        for cid, stt in states.items():
            ip = cams.get(cid) or ""
            sd = bool(ip and ".".join(ip.split(".")[:3]) in dark and stt != "recording")
            prev = _last_rows.get(cid)
            if prev and prev[1] == stt and prev[2] == sd:
                keep.append(prev[0])
                continue
            if prev:
                cur.execute("UPDATE mes_vcov_cam_state SET to_ts = now() WHERE id = %s",
                            (prev[0],))
            cur.execute("""INSERT INTO mes_vcov_cam_state
                               (camera_id, ip, state, subnet_down, from_ts, to_ts)
                           VALUES (%s, %s, %s, %s, now(), now()) RETURNING id""",
                        (cid, ip, stt, sd))
            _last_rows[cid] = (cur.fetchone()[0], stt, sd)
        if keep:
            cur.execute("UPDATE mes_vcov_cam_state SET to_ts = now() WHERE id = ANY(%s)",
                        (keep,))
    _agent_mark(last_sample=datetime.now().astimezone())


# ── cycle evaluator ─────────────────────────────────────────────────────────
def _clip_dir_cache():
    cache = {}

    def has(kind, owner, rec_date, seq, ng, shift, line_id):
        from routers.clip_archive import clip_path
        p = clip_path(kind, owner, rec_date, seq, ng, shift, line_id)
        d, leaf = os.path.split(p)
        if d not in cache:
            try:
                cache[d] = set(os.listdir(d))
            except Exception:
                cache[d] = set()
        other = leaf.replace("_ng.mp4", ".mp4") if leaf.endswith("_ng.mp4") \
            else leaf.replace(".mp4", "_ng.mp4")
        return leaf in cache[d] or other in cache[d]
    return has


def _load_states(cur, t0, t1):
    cur.execute("""SELECT camera_id, state, subnet_down, from_ts, to_ts
                     FROM mes_vcov_cam_state
                    WHERE to_ts >= %s AND from_ts <= %s""",
                (t0 - timedelta(seconds=SLACK_S), t1 + timedelta(seconds=SLACK_S)))
    out = {}
    for r in cur.fetchall():
        out.setdefault(r[0], []).append(r[1:])
    return out


def _near_wipe(ts_end, bounds):
    """True if the cycle ended in the last 42 min before a CMS shift start —
    the footage its clip needed was deleted by the shift-wipe."""
    local = ts_end.astimezone() if ts_end.tzinfo else ts_end
    for day in (local.date(), local.date() + timedelta(days=1)):
        for h, m in bounds:
            b = datetime(day.year, day.month, day.day, h, m, tzinfo=local.tzinfo)
            if timedelta(0) < b - local <= timedelta(minutes=42):
                return True
    return False


def _reason(cam, ts_start, ts_end, states, bounds):
    if not cam:
        return "no_camera", "no camera bound to this machine"
    a = (ts_start or ts_end) - timedelta(seconds=5)
    b = ts_end + timedelta(seconds=5)
    hit = [r for r in states.get(cam, [])
           if r[2] - timedelta(seconds=SLACK_S) <= b and r[3] + timedelta(seconds=SLACK_S) >= a]
    if not hit:
        return "unknown", "tracker has no sample for this time"
    worst = min(hit, key=lambda r: _STATE_RANK.get(r[0], 5))
    stt, subnet = worst[0], worst[1]
    if stt == "cms_down":
        return "cms_down", "CMS service was not answering"
    if stt in ("camera_offline", "camera_hung") and subnet:
        return "network_down", f"{stt.replace('_', ' ')}; 80%+ of cameras on its subnet down"
    if stt == "camera_offline":
        return "camera_offline", "no ping reply, no video file growing"
    if stt == "camera_hung":
        return "camera_hung", "answers ping but its video file was not growing"
    if _near_wipe(ts_end, bounds):
        return "shift_wipe", "ended <42 min before a CMS shift start (footage wiped)"
    return "clip_failed", "camera was recording but no clip was cut in time"


def _naive(dt):
    """Local wall-clock time without tzinfo — main-line *_ct_log.ts is a plain
    `timestamp` in plant time, so it is compared naive to keep its index."""
    return dt.astimezone().replace(tzinfo=None) if dt.tzinfo else dt


def _cursor(cur, source, now):
    cur.execute("SELECT last_ts FROM mes_vcov_cursor WHERE source = %s", (source,))
    r = cur.fetchone()
    if r:
        return r["last_ts"]
    # go-live: judge only cycles from now on (operator chose no backfill)
    cur.execute("INSERT INTO mes_vcov_cursor (source, last_ts) VALUES (%s, %s) "
                "ON CONFLICT (source) DO NOTHING", (source, now))
    return now


def _evaluate():
    _ensure_tables()
    cams_ip = _load_cameras()
    bounds = _shift_boundaries()
    has_clip = _clip_dir_cache()
    with _lock, get_conn() as conn:
        cur = dict_cursor(conn)
        cur.execute("SELECT now() AS n")
        now = cur.fetchone()["n"]
        horizon = now - timedelta(minutes=MATURE_MIN)
        lines, subs = _machines(cur)
        line_by_id = {ln["id"]: ln for ln in lines}
        cycles = []            # (line_id, machine_key, machine_name, cam, row)
        cursors = {}

        for ln in lines:
            src = f"line:{ln['id']}"
            last = _cursor(cur, src, now)
            if last >= horizon:
                continue
            tbl = f"{ln['db_table_name']}_ct_log"
            if not _TABLE_RE.match(tbl):
                continue
            try:
                cur.execute("SAVEPOINT vcov_line")
                cur.execute(f"""SELECT cycle_seq, record_date, shift_name,
                                       ts::timestamptz AS ts_end,
                                       ct_value AS ct, part_code,
                                       COALESCE(is_ng, FALSE) AS is_ng
                                  FROM {tbl}
                                 WHERE ts > %s AND ts <= %s
                                 ORDER BY ts LIMIT 20000""", (_naive(last), _naive(horizon)))
                rows = cur.fetchall()
                cur.execute("RELEASE SAVEPOINT vcov_line")
            except Exception:
                cur.execute("ROLLBACK TO SAVEPOINT vcov_line")
                continue
            cursors[src] = rows[-1]["ts_end"] if len(rows) == 20000 else horizon
            for r in rows:
                cycles.append((ln["id"], "main", ln.get("mname") or "Final Inspection",
                               ln.get("cam"), r, "line", ln["id"]))

        last = _cursor(cur, "subs", now)
        if last < horizon:
            cur.execute("""SELECT sub_plc_id, line_id, cycle_seq, record_date, shift_name,
                                  ts_start, ts_end, ct_seconds AS ct, part_code,
                                  COALESCE(is_ng, FALSE) AS is_ng
                             FROM mes_submachine_ct_log
                            WHERE record_date >= %s AND ts_end > %s AND ts_end <= %s
                            ORDER BY ts_end LIMIT 60000""",
                        ((last - timedelta(days=1)).date(), last, horizon))
            rows = cur.fetchall()
            cursors["subs"] = rows[-1]["ts_end"] if len(rows) == 60000 else horizon
            for r in rows:
                sp = subs.get(r["sub_plc_id"]) or {}
                lid = r["line_id"] or sp.get("line_id")
                if lid not in line_by_id:
                    continue
                cycles.append((lid, f"sub_{r['sub_plc_id']}",
                               sp.get("mname") or f"Sub {r['sub_plc_id']}",
                               sp.get("cam"), r, "sub", r["sub_plc_id"]))

        if cycles:
            t0 = min((c[4].get("ts_start") or c[4]["ts_end"]) for c in cycles)
            t1 = max(c[4]["ts_end"] for c in cycles)
            states = _load_states(conn.cursor(), t0, t1)
        else:
            states = {}

        hourly, missing, seen = {}, [], set()
        for lid, mkey, mname, cam, r, kind, owner in cycles:
            ck = (lid, mkey, r["record_date"], r.get("shift_name"), r["cycle_seq"], r["is_ng"])
            if ck in seen:
                continue          # a repeated log row is the same cycle and the same clip
            seen.add(ck)
            ts_end = r["ts_end"]
            ct = float(r["ct"]) if r.get("ct") is not None else None
            ts_start = r.get("ts_start") or (ts_end - timedelta(seconds=ct) if ct else ts_end)
            shift = r.get("shift_name") or "UNKNOWN"
            ok = has_clip(kind, owner, r["record_date"], r["cycle_seq"], r["is_ng"], shift, lid)
            hk = (r["record_date"], shift, lid, mkey,
                  ts_end.replace(minute=0, second=0, microsecond=0))
            h = hourly.setdefault(hk, [mname, 0, 0, 0])
            h[1] += 1
            if ok:
                h[2] += 1
                continue
            h[3] += 1
            reason, detail = _reason(cam, ts_start, ts_end, states, bounds)
            missing.append((r["record_date"], shift, lid, mkey, mname, r["cycle_seq"],
                            r["is_ng"], ts_start, ts_end, r.get("part_code"), ct,
                            cam, cams_ip.get(cam or "", ""), reason, detail))

        pc = conn.cursor()
        for (rd, sh, lid, mkey, hr), (mname, n, c, m) in hourly.items():
            pc.execute("""INSERT INTO mes_vcov_hourly
                              (record_date, shift_name, line_id, machine_key, machine_name,
                               hour_ts, cycles, clips, missing)
                          VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)
                          ON CONFLICT (record_date, shift_name, line_id, machine_key, hour_ts)
                          DO UPDATE SET cycles = mes_vcov_hourly.cycles + EXCLUDED.cycles,
                                        clips = mes_vcov_hourly.clips + EXCLUDED.clips,
                                        missing = mes_vcov_hourly.missing + EXCLUDED.missing,
                                        machine_name = EXCLUDED.machine_name""",
                       (rd, sh, lid, mkey, mname, hr, n, c, m))
        for row in missing:
            pc.execute("""INSERT INTO mes_vcov_missing
                              (record_date, shift_name, line_id, machine_key, machine_name,
                               cycle_seq, is_ng, ts_start, ts_end, part_code, ct,
                               camera_id, camera_ip, reason, detail)
                          VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                          ON CONFLICT (line_id, machine_key, record_date, shift_name,
                                       cycle_seq, is_ng) DO NOTHING""", row)
        for src, ts in cursors.items():
            pc.execute("UPDATE mes_vcov_cursor SET last_ts = %s WHERE source = %s", (ts, src))

        # a clip that arrived late (archiver backlog) turns a "missing" into a clip
        cur.execute("""SELECT id, record_date, shift_name, line_id, machine_key,
                              cycle_seq, is_ng, ts_end
                         FROM mes_vcov_missing
                        WHERE ts_end > now() - interval '%s hours'
                          AND reason IN ('clip_failed', 'shift_wipe', 'unknown',
                                         'camera_hung', 'camera_offline',
                                         'network_down', 'cms_down')""" % RECHECK_H)
        for r in cur.fetchall():
            kind, owner = (("line", r["line_id"]) if r["machine_key"] == "main"
                           else ("sub", int(r["machine_key"][4:])))
            if has_clip(kind, owner, r["record_date"], r["cycle_seq"], r["is_ng"],
                        r["shift_name"], r["line_id"]):
                pc.execute("DELETE FROM mes_vcov_missing WHERE id = %s", (r["id"],))
                pc.execute("""UPDATE mes_vcov_hourly
                                 SET clips = clips + 1, missing = GREATEST(missing - 1, 0)
                               WHERE record_date = %s AND shift_name = %s AND line_id = %s
                                 AND machine_key = %s AND hour_ts = %s""",
                           (r["record_date"], r["shift_name"], r["line_id"],
                            r["machine_key"],
                            r["ts_end"].replace(minute=0, second=0, microsecond=0)))
    _agent_mark(last_eval=datetime.now().astimezone())
    return len(cycles), len(missing)


def _retention():
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute("DELETE FROM mes_vcov_missing WHERE ts_end < now() - interval '%s days'"
                    % RETAIN_DAYS)
        cur.execute("DELETE FROM mes_vcov_hourly WHERE hour_ts < now() - interval '%s days'"
                    % RETAIN_DAYS)
        cur.execute("DELETE FROM mes_vcov_cam_state WHERE to_ts < now() - interval '%s days'"
                    % RETAIN_DAYS)
        cur.execute("DELETE FROM mes_vcov_findings WHERE closed_at < now() - interval '%s days'"
                    % RETAIN_DAYS)


# ── the Video Agent — watches and reports, never acts ───────────────────────
def _agent():
    """Open / close findings and send one digest to admins.  READ-ONLY by
    design: there is deliberately no restart, kill or config code here."""
    _ensure_tables()
    found = {}      # fkey -> (kind, line_id, camera_id, message)
    with get_conn() as conn:
        cur = dict_cursor(conn)
        lines, subs = _machines(cur)
        lname = {ln["id"]: ln["line_name"] for ln in lines}
        # CMS down for 2+ consecutive samples
        cur.execute("""SELECT count(*) AS n FROM mes_vcov_cam_state
                        WHERE state = 'cms_down' AND to_ts > now() - interval '3 minutes'
                          AND to_ts - from_ts >= interval '110 seconds'""")
        if cur.fetchone()["n"]:
            found["cms"] = ("CMS_DOWN", None, None, "CMS is not answering — no camera is recording")

        # cameras dark for 10+ min whose machine produced in the last 15 min
        needed = {}
        for ln in lines:
            tbl = f"{ln['db_table_name']}_ct_log"
            if not ln.get("cam") or not _TABLE_RE.match(tbl):
                continue
            try:
                cur.execute("SAVEPOINT vcov_ag")
                cur.execute(f"SELECT 1 FROM {tbl} WHERE ts > now() - interval '15 minutes' LIMIT 1")
                if cur.fetchone():
                    needed.setdefault(ln["cam"], (ln["id"], ln.get("mname") or "Final Inspection"))
                cur.execute("RELEASE SAVEPOINT vcov_ag")
            except Exception:
                cur.execute("ROLLBACK TO SAVEPOINT vcov_ag")
        cur.execute("""SELECT DISTINCT sub_plc_id FROM mes_submachine_ct_log
                        WHERE record_date >= current_date - 1
                          AND ts_end > now() - interval '15 minutes'""")
        for r in cur.fetchall():
            sp = subs.get(r["sub_plc_id"]) or {}
            if sp.get("cam"):
                needed.setdefault(sp["cam"], (sp.get("line_id"), sp.get("mname") or "Sub"))
        if needed:
            cur.execute("""SELECT DISTINCT ON (camera_id) camera_id, ip, state, from_ts
                             FROM mes_vcov_cam_state
                            WHERE camera_id = ANY(%s) AND to_ts > now() - interval '3 minutes'
                            ORDER BY camera_id, to_ts DESC""", (list(needed),))
            for r in cur.fetchall():
                if r["state"] in ("camera_offline", "camera_hung") and \
                        r["from_ts"] <= datetime.now().astimezone() - timedelta(minutes=10):
                    lid, mname = needed[r["camera_id"]]
                    why = "offline (no ping)" if r["state"] == "camera_offline" else "hung (ping OK, no video)"
                    found[f"cam:{r['camera_id']}"] = (
                        "CAMERA_DOWN", lid, r["camera_id"],
                        f"{lname.get(lid, lid)} · {mname} camera {r['ip'] or ''} {why} "
                        f"since {r['from_ts'].astimezone():%H:%M}")

        # line coverage over the last judged hour
        cur.execute("""SELECT line_id, sum(cycles) AS n, sum(clips) AS c
                         FROM mes_vcov_hourly
                        WHERE hour_ts >= now() - interval '%s minutes'
                          AND hour_ts <  now() - interval '%s minutes'
                        GROUP BY line_id""" % (MATURE_MIN + 75, MATURE_MIN))
        cov = {r["line_id"]: (int(r["n"] or 0), int(r["c"] or 0)) for r in cur.fetchall()}
        cur.execute("SELECT fkey FROM mes_vcov_findings WHERE closed_at IS NULL AND kind = 'LOW_COVERAGE'")
        open_cov = {r["fkey"] for r in cur.fetchall()}
        for lid, (n, c) in cov.items():
            pct = c * 100.0 / n if n else 100.0
            key = f"cov:{lid}"
            if n >= 20 and (pct < 80 or (key in open_cov and pct < 90)):
                found[key] = ("LOW_COVERAGE", lid, None,
                              f"{lname.get(lid, lid)} video coverage {pct:.0f}% ({c}/{n}) in the last judged hour")

        cur.execute("SELECT id, fkey FROM mes_vcov_findings WHERE closed_at IS NULL")
        open_now = {r["fkey"]: r["id"] for r in cur.fetchall()}
        pc = conn.cursor()
        opened, closed = [], []
        for key, (kind, lid, cam, msg) in found.items():
            if key in open_now:
                pc.execute("UPDATE mes_vcov_findings SET last_seen = now(), message = %s WHERE id = %s",
                           (msg, open_now[key]))
            else:
                pc.execute("""INSERT INTO mes_vcov_findings (kind, fkey, line_id, camera_id, message)
                              VALUES (%s,%s,%s,%s,%s)""", (kind, key, lid, cam, msg))
                opened.append(msg)
        for key, fid in open_now.items():
            if key not in found:
                pc.execute("UPDATE mes_vcov_findings SET closed_at = now() WHERE id = %s", (fid,))
                closed.append(key)
    _agent_mark(last_agent=datetime.now().astimezone())
    if opened or closed:
        _notify(opened, closed)


_pending_note = {"opened": [], "closed": []}


def _notify(opened, closed):
    """One Inbox digest to admins, at most every 15 minutes."""
    _pending_note["opened"].extend(opened)
    _pending_note["closed"].extend(closed)
    try:
        with get_conn() as conn:
            cur = dict_cursor(conn)
            cur.execute("SELECT last_notify FROM mes_vcov_agent WHERE id = 1")
            last = (cur.fetchone() or {}).get("last_notify")
            if last and (datetime.now().astimezone() - last).total_seconds() < NOTIFY_MIN_GAP_S:
                return
            cur.execute("SELECT id FROM users WHERE role = 'admin' AND COALESCE(is_active, TRUE)")
            admins = [r["id"] for r in cur.fetchall()]
    except Exception:
        return
    op, cl = _pending_note["opened"], _pending_note["closed"]
    if not (op or cl):
        return
    if op:
        title = f"Video Agent: {len(op)} new issue{'s' if len(op) != 1 else ''}" + \
                (f", {len(cl)} resolved" if cl else "")
    else:
        title = f"Video Agent: {len(cl)} issue{'s' if len(cl) != 1 else ''} resolved"
    body = "; ".join(op[:6]) + (f" … +{len(op) - 6} more" if len(op) > 6 else "")
    if not op:
        body = f"{len(cl)} video issue(s) cleared."
    try:
        from routers.push import send_to_user
        for uid in admins:
            send_to_user(uid, title, body, url="/video-coverage", tag="video-agent")
        _agent_mark(last_notify=datetime.now().astimezone())
        _pending_note["opened"], _pending_note["closed"] = [], []
    except Exception as exc:
        print(f"[VCOV] notify failed: {exc}", flush=True)


# ── Camera Status (live, zone → line → camera) + its logs ─────────────────
# Zone / line / camera: cameras online, clips being cut or not, and why, plus a
# camera log and a 5-minute line log.  LIVE: cycles that ended 3–20 min ago against
# the clip archive (the archiver cuts within a couple of minutes), so this does
# not wait the 45-min judging delay of the coverage summary.
LIVE_FROM_MIN, LIVE_TO_MIN = 20, 3
NOTCUT_REASON = {
    "cms_down":       "CMS down",
    "network_down":   "Network / switch down",
    "camera_offline": "Camera offline (no ping)",
    "camera_hung":    "Camera hung (ping OK, no video)",
    "recording":      "Recording, but clips not being cut",
    None:             "Not tracked yet",
}


@once
def _ensure_status_tables():
    with get_conn() as conn:
        cur = conn.cursor()
        cur.execute("""
            CREATE TABLE IF NOT EXISTS mes_vcov_line_snap (
                ts          TIMESTAMPTZ NOT NULL,
                line_id     INT NOT NULL,
                zone_name   TEXT,
                line_name   TEXT,
                cameras     INT, online INT, hung INT, offline INT, cms_down INT,
                cutting     INT, not_cutting INT, idle INT,
                PRIMARY KEY (ts, line_id))""")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_vcov_line_snap_line "
                    "ON mes_vcov_line_snap (line_id, ts)")


def _camera_status(cur, allowed=None, include_unassigned=False):
    """Every bound camera with its live state and whether its clips are being
    cut, grouped zone → line.  `allowed` = line ids the caller may see."""
    lines, subs = _machines(cur)
    lmeta = {ln["id"]: ln for ln in lines}
    cams_ip = _load_cameras()
    by_cam = {}                                  # cam -> list of machine dicts
    for ln in lines:
        if ln.get("cam"):
            by_cam.setdefault(ln["cam"], []).append(
                {"kind": "line", "owner": ln["id"], "line_id": ln["id"], "key": "main",
                 "name": ln.get("mname") or "Final Inspection"})
    for sid, sp in subs.items():
        if sp.get("cam") and sp.get("line_id") in lmeta:
            by_cam.setdefault(sp["cam"], []).append(
                {"kind": "sub", "owner": sid, "line_id": sp["line_id"], "key": f"sub_{sid}",
                 "name": sp.get("mname") or f"Sub {sid}"})
    cur.execute("""SELECT DISTINCT ON (camera_id) camera_id, state, subnet_down, from_ts, to_ts
                     FROM mes_vcov_cam_state
                    WHERE to_ts > now() - interval '3 minutes'
                    ORDER BY camera_id, to_ts DESC""")
    st = {r["camera_id"]: r for r in cur.fetchall()}
    cur.execute("SELECT now() AS n")
    now = cur.fetchone()["n"]
    t0, t1 = now - timedelta(minutes=LIVE_FROM_MIN), now - timedelta(minutes=LIVE_TO_MIN)
    has_clip = _clip_dir_cache()
    counts = {}                                  # (kind, owner) -> [cycles, clips]
    seen = set()
    for ln in lines:
        tbl = f"{ln['db_table_name']}_ct_log"
        if not ln.get("cam") or not _TABLE_RE.match(tbl):
            continue
        try:
            cur.execute("SAVEPOINT vcov_live")
            cur.execute(f"""SELECT cycle_seq, record_date, shift_name, COALESCE(is_ng, FALSE) AS is_ng
                              FROM {tbl} WHERE ts > %s AND ts <= %s""", (_naive(t0), _naive(t1)))
            rows = cur.fetchall()
            cur.execute("RELEASE SAVEPOINT vcov_live")
        except Exception:
            cur.execute("ROLLBACK TO SAVEPOINT vcov_live")
            continue
        for r in rows:
            k = ("line", ln["id"], r["record_date"], r["shift_name"], r["cycle_seq"], r["is_ng"])
            if k in seen:
                continue
            seen.add(k)
            c = counts.setdefault(("line", ln["id"]), [0, 0])
            c[0] += 1
            c[1] += has_clip("line", ln["id"], r["record_date"], r["cycle_seq"], r["is_ng"],
                             r["shift_name"] or "UNKNOWN", ln["id"])
    cur.execute("""SELECT sub_plc_id, line_id, cycle_seq, record_date, shift_name,
                          COALESCE(is_ng, FALSE) AS is_ng
                     FROM mes_submachine_ct_log
                    WHERE record_date >= current_date - 1 AND ts_end > %s AND ts_end <= %s""", (t0, t1))
    for r in cur.fetchall():
        sp = subs.get(r["sub_plc_id"]) or {}
        lid = r["line_id"] or sp.get("line_id")
        if not sp.get("cam") or lid not in lmeta:
            continue
        k = ("sub", r["sub_plc_id"], r["record_date"], r["shift_name"], r["cycle_seq"], r["is_ng"])
        if k in seen:
            continue
        seen.add(k)
        c = counts.setdefault(("sub", r["sub_plc_id"]), [0, 0])
        c[0] += 1
        c[1] += has_clip("sub", r["sub_plc_id"], r["record_date"], r["cycle_seq"], r["is_ng"],
                         r["shift_name"] or "UNKNOWN", lid)

    zones = {}
    for cam, machines in by_cam.items():
        lid = machines[0]["line_id"]
        if allowed is not None and lid not in allowed:
            continue
        s_ = st.get(cam)
        state = s_["state"] if s_ else None
        n = sum(counts.get((m["kind"], m["owner"]), [0, 0])[0] for m in machines)
        c = sum(counts.get((m["kind"], m["owner"]), [0, 0])[1] for m in machines)
        # Cutting = this camera's clips ARE being produced.  The archiver cuts
        # Final Inspection first and fills sub-machine clips in over its 42-min
        # window, so a partial share here is normally "catching up", not a fault;
        # only zero clips for a producing machine is "not cutting".
        if n == 0:
            status, reason = "idle", "Machine made no cycles in the window"
        elif c > 0:
            status = "cutting"
            reason = None if c * 100 >= 80 * n else f"Clips catching up ({round(c * 100.0 / n)}% so far)"
        else:
            key = "network_down" if (s_ and s_["subnet_down"] and state != "recording") else state
            status, reason = "not_cutting", NOTCUT_REASON.get(key, NOTCUT_REASON[None])
        ln = lmeta[lid]
        z = zones.setdefault(ln.get("zone_name") or "—", {})
        L = z.setdefault(lid, {"line_id": lid, "line_name": ln["line_name"],
                               "zone_name": ln.get("zone_name") or "—", "cameras": []})
        L["cameras"].append({
            "camera_id": cam, "ip": cams_ip.get(cam, ""),
            "machines": [m["name"] for m in machines if m["line_id"] == lid] or [m["name"] for m in machines],
            "main": any(m["key"] == "main" for m in machines),
            "state": state, "subnet_down": bool(s_ and s_["subnet_down"]),
            "since": s_["from_ts"].astimezone().isoformat() if s_ else None,
            "cycles": n, "clips": c, "clip_pct": round(c * 100.0 / n, 1) if n else None,
            "status": status, "reason": reason})
    out, tot = [], dict(cameras=0, online=0, hung=0, offline=0, cms_down=0,
                        cutting=0, not_cutting=0, idle=0)
    for zname in sorted(zones):
        zl = []
        for L in sorted(zones[zname].values(), key=lambda x: x["line_name"]):
            cs = L["cameras"]
            cnt = dict(cameras=len(cs),
                       online=sum(1 for x in cs if x["state"] == "recording"),
                       hung=sum(1 for x in cs if x["state"] == "camera_hung"),
                       offline=sum(1 for x in cs if x["state"] == "camera_offline"),
                       cms_down=sum(1 for x in cs if x["state"] == "cms_down"),
                       cutting=sum(1 for x in cs if x["status"] == "cutting"),
                       not_cutting=sum(1 for x in cs if x["status"] == "not_cutting"),
                       idle=sum(1 for x in cs if x["status"] == "idle"))
            for k2, v in cnt.items():
                tot[k2] += v
            order = {"not_cutting": 0, "cutting": 1, "idle": 2}
            cs.sort(key=lambda x: (not x["main"], order.get(x["status"], 3), x["ip"]))
            zl.append({**{k2: L[k2] for k2 in ("line_id", "line_name", "zone_name")}, **cnt,
                       "cameras_list": cs})
        out.append({"zone_name": zname, "lines": zl,
                    **{k2: sum(l[k2] for l in zl) for k2 in tot}})
    tot["total_cameras"] = len(cams_ip)
    unbound = sorted(set(cams_ip) - set(by_cam))
    tot["unbound"] = len(unbound)
    unassigned = []
    if include_unassigned:
        for cam in unbound:
            s_ = st.get(cam)
            unassigned.append({"camera_id": cam, "ip": cams_ip.get(cam, ""),
                               "state": s_["state"] if s_ else None,
                               "subnet_down": bool(s_ and s_["subnet_down"]),
                               "since": s_["from_ts"].astimezone().isoformat() if s_ else None})
    return {"window": {"from": t0.astimezone().isoformat(), "to": t1.astimezone().isoformat()},
            "kpis": tot, "zones": out, "unassigned": unassigned}


def _snapshot():
    """Every 5 min: each line's camera counts, kept as the line log."""
    _ensure_status_tables()
    with get_conn() as conn:
        cur = dict_cursor(conn)
        data = _camera_status(cur)
        pc = conn.cursor()
        for z in data["zones"]:
            for L in z["lines"]:
                pc.execute("""INSERT INTO mes_vcov_line_snap
                                  (ts, line_id, zone_name, line_name, cameras, online, hung,
                                   offline, cms_down, cutting, not_cutting, idle)
                              VALUES (date_trunc('minute', now()), %s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                              ON CONFLICT (ts, line_id) DO NOTHING""",
                           (L["line_id"], L["zone_name"], L["line_name"], L["cameras"], L["online"],
                            L["hung"], L["offline"], L["cms_down"], L["cutting"], L["not_cutting"],
                            L["idle"]))
        pc.execute("DELETE FROM mes_vcov_line_snap WHERE ts < now() - interval '%s days'" % RETAIN_DAYS)


# ── worker loop (runs in the ONE background-leader process) ────────────────
def _loop():
    try:
        _ensure_tables()
        now = datetime.now().astimezone()
        with get_conn() as conn:
            cur = conn.cursor()
            cur.execute("""UPDATE mes_vcov_agent SET started_at = %s,
                                  tracking_since = COALESCE(tracking_since, %s)
                            WHERE id = 1""", (now, now))
    except Exception as exc:
        print(f"[VCOV] start failed: {exc}", flush=True)
    nxt = {"sample": 0.0, "eval": time.time() + 90, "agent": time.time() + 150,
           "snap": time.time() + 75, "ret": 0.0}
    while True:
        t = time.time()
        for name, fn, every in (("sample", _sample, SAMPLE_S), ("eval", _evaluate, EVAL_S),
                                ("agent", _agent, AGENT_S), ("snap", _snapshot, 300),
                                ("ret", _retention, 6 * 3600)):
            if t >= nxt[name]:
                nxt[name] = t + every
                try:
                    fn()
                except Exception as exc:
                    print(f"[VCOV] {name} failed: {exc}", flush=True)
                    _agent_mark(last_error=f"{datetime.now():%H:%M:%S} {name}: {exc}"[:500])
        time.sleep(5)


_STARTED = False


def _start():
    global _STARTED
    if _STARTED or os.environ.get("VCOV_TRACKER", "1") == "0":
        return
    _STARTED = True
    threading.Thread(target=_loop, name="video-coverage", daemon=True).start()
    print("[VCOV] Video Coverage tracker + Video Agent started (read-only)", flush=True)


from bg_leader import is_leader as _bg_is_leader   # noqa: E402
if _bg_is_leader():
    _start()


# ── API ─────────────────────────────────────────────────────────────────────
def _scope(cur, user, line_id):
    from routers.shift_compile import _accessible_lines
    allowed = {r["id"] for r in _accessible_lines(cur, user)}
    if line_id is not None:
        if line_id not in allowed:
            raise HTTPException(403, "Not authorized for this line")
        return [line_id]
    return sorted(allowed)


def _range(date_from, date_to):
    try:
        d0 = date.fromisoformat(date_from) if date_from else date.today()
        d1 = date.fromisoformat(date_to) if date_to else d0
    except ValueError:
        raise HTTPException(400, "dates must be YYYY-MM-DD")
    if d1 < d0:
        d0, d1 = d1, d0
    if (d1 - d0).days > 31:
        raise HTTPException(400, "range too large (max 31 days)")
    return d0, d1


def _where(d0, d1, shift, lines, prefix=""):
    w = [f"{prefix}record_date BETWEEN %s AND %s", f"{prefix}line_id = ANY(%s)"]
    a = [d0, d1, lines]
    if shift and shift != "ALL":
        w.append(f"{prefix}shift_name = %s")
        a.append(shift)
    return " AND ".join(w), a


@router.get("/summary")
def summary(date_from: Optional[str] = Query(None), date_to: Optional[str] = Query(None),
            shift: str = Query("ALL"), line_id: Optional[int] = Query(None),
            user=Depends(get_current_user)):
    """KPIs, per line, per machine (when a line is picked), per reason, hourly."""
    _ensure_tables()
    d0, d1 = _range(date_from, date_to)
    with get_conn() as conn:
        cur = dict_cursor(conn)
        lines = _scope(cur, user, line_id)
        cur.execute("SELECT id, line_name, zone_id, (SELECT zone_name FROM mes_zones z WHERE z.id = l.zone_id) AS zone_name "
                    "FROM mes_lines l WHERE id = ANY(%s)", (lines,))
        meta = {r["id"]: r for r in cur.fetchall()}
        w, a = _where(d0, d1, shift, lines)
        cur.execute(f"""SELECT line_id, sum(cycles) n, sum(clips) c, sum(missing) m
                          FROM mes_vcov_hourly WHERE {w} GROUP BY line_id""", a)
        per_line = {r["line_id"]: r for r in cur.fetchall()}
        cur.execute(f"""SELECT line_id, reason, count(*) n FROM mes_vcov_missing
                         WHERE {w} GROUP BY line_id, reason""", a)
        reasons_line = {}
        reason_tot = {}
        for r in cur.fetchall():
            reasons_line.setdefault(r["line_id"], {})[r["reason"]] = r["n"]
            reason_tot[r["reason"]] = reason_tot.get(r["reason"], 0) + r["n"]
        cur.execute(f"""SELECT line_id, date_trunc('hour', hour_ts) h, sum(cycles) n, sum(clips) c
                          FROM mes_vcov_hourly WHERE {w}
                         GROUP BY line_id, h ORDER BY h""", a)
        hourly = [{"line_id": r["line_id"], "hour": r["h"].astimezone().isoformat(),
                   "cycles": int(r["n"]), "clips": int(r["c"])} for r in cur.fetchall()]
        machines = []
        if line_id is not None:
            cur.execute(f"""SELECT machine_key, max(machine_name) mname, sum(cycles) n,
                                   sum(clips) c, sum(missing) m
                              FROM mes_vcov_hourly WHERE {w}
                             GROUP BY machine_key ORDER BY machine_key = 'main' DESC, machine_key""", a)
            mrows = cur.fetchall()
            cur.execute(f"""SELECT machine_key, reason, count(*) n, max(camera_id) cam, max(camera_ip) ip
                              FROM mes_vcov_missing WHERE {w}
                             GROUP BY machine_key, reason""", a)
            mreason, mcam = {}, {}
            for r in cur.fetchall():
                mreason.setdefault(r["machine_key"], {})[r["reason"]] = r["n"]
                if r["cam"]:
                    mcam[r["machine_key"]] = (r["cam"], r["ip"])
            for r in mrows:
                n, c = int(r["n"] or 0), int(r["c"] or 0)
                cam = mcam.get(r["machine_key"], (None, None))
                machines.append({"machine_key": r["machine_key"], "machine_name": r["mname"],
                                 "camera_id": cam[0], "camera_ip": cam[1],
                                 "cycles": n, "clips": c, "missing": int(r["m"] or 0),
                                 "coverage": round(c * 100.0 / n, 1) if n else None,
                                 "reasons": mreason.get(r["machine_key"], {})})
        cur.execute("SELECT tracking_since, last_eval FROM mes_vcov_agent WHERE id = 1")
        ag = cur.fetchone() or {}
        pending = 0
        if d1 >= date.today():
            try:
                cur.execute(f"""SELECT count(*) n FROM mes_submachine_ct_log
                                 WHERE record_date >= current_date - 1 AND line_id = ANY(%s)
                                   AND ts_end > now() - interval '{MATURE_MIN} minutes'""",
                            (lines,))
                pending = int(cur.fetchone()["n"])
                cur.execute("SELECT db_table_name FROM mes_lines WHERE id = ANY(%s)", (lines,))
                for t in [r["db_table_name"] for r in cur.fetchall()]:
                    if not t or not _TABLE_RE.match(f"{t}_ct_log"):
                        continue
                    cur.execute("SAVEPOINT vcov_pend")
                    try:
                        cur.execute(f"""SELECT count(*) n FROM {t}_ct_log
                                         WHERE ts > (now() - interval '{MATURE_MIN} minutes')::timestamp""")
                        pending += int(cur.fetchone()["n"])
                        cur.execute("RELEASE SAVEPOINT vcov_pend")
                    except Exception:
                        cur.execute("ROLLBACK TO SAVEPOINT vcov_pend")
            except Exception:
                conn.rollback()
    out_lines, tot = [], {"cycles": 0, "clips": 0, "missing": 0}
    for lid in lines:
        r = per_line.get(lid)
        if not r:
            continue
        n, c, m = int(r["n"] or 0), int(r["c"] or 0), int(r["m"] or 0)
        tot["cycles"] += n; tot["clips"] += c; tot["missing"] += m
        mt = meta.get(lid, {})
        out_lines.append({"line_id": lid, "line_name": mt.get("line_name"),
                          "zone_name": mt.get("zone_name"), "cycles": n, "clips": c,
                          "missing": m, "coverage": round(c * 100.0 / n, 1) if n else None,
                          "reasons": reasons_line.get(lid, {})})
    out_lines.sort(key=lambda x: (x["coverage"] if x["coverage"] is not None else 101))
    tot["coverage"] = round(tot["clips"] * 100.0 / tot["cycles"], 1) if tot["cycles"] else None
    tot["pending"] = pending
    return {"date_from": d0.isoformat(), "date_to": d1.isoformat(), "shift": shift,
            "tracking_since": ag.get("tracking_since").isoformat() if ag.get("tracking_since") else None,
            "last_eval": ag.get("last_eval").isoformat() if ag.get("last_eval") else None,
            "mature_min": MATURE_MIN, "kpis": tot, "lines": out_lines, "machines": machines,
            "reasons": sorted(({"key": k, "label": REASONS.get(k, k), "count": v}
                               for k, v in reason_tot.items()), key=lambda x: -x["count"]),
            "reason_labels": REASONS, "hourly": hourly}


def _missing_rows(cur, d0, d1, shift, lines, machine_key, reason, limit, offset):
    w, a = _where(d0, d1, shift, lines, "m.")
    if machine_key:
        w += " AND m.machine_key = %s"
        a.append(machine_key)
    if reason:
        w += " AND m.reason = %s"
        a.append(reason)
    cur.execute(f"SELECT count(*) n FROM mes_vcov_missing m WHERE {w}", a)
    total = int(cur.fetchone()["n"])
    cur.execute(f"""SELECT m.*, l.line_name
                      FROM mes_vcov_missing m JOIN mes_lines l ON l.id = m.line_id
                     WHERE {w} ORDER BY m.ts_end DESC LIMIT %s OFFSET %s""",
                a + [limit, offset])
    rows = []
    for r in cur.fetchall():
        rows.append({
            "record_date": r["record_date"].isoformat(), "shift": r["shift_name"],
            "line_id": r["line_id"], "line_name": r["line_name"],
            "machine_key": r["machine_key"], "machine_name": r["machine_name"],
            "cycle_seq": r["cycle_seq"], "is_ng": r["is_ng"], "part_code": r["part_code"],
            "ts_start": r["ts_start"].astimezone().isoformat() if r["ts_start"] else None,
            "ts_end": r["ts_end"].astimezone().isoformat(),
            "ct": float(r["ct"]) if r["ct"] is not None else None,
            "camera_id": r["camera_id"], "camera_ip": r["camera_ip"],
            "reason": r["reason"], "reason_label": REASONS.get(r["reason"], r["reason"]),
            "detail": r["detail"]})
    return total, rows


@router.get("/missing")
def missing(date_from: Optional[str] = Query(None), date_to: Optional[str] = Query(None),
            shift: str = Query("ALL"), line_id: Optional[int] = Query(None),
            machine_key: Optional[str] = Query(None), reason: Optional[str] = Query(None),
            page: int = Query(1, ge=1), page_size: int = Query(100, ge=10, le=500),
            user=Depends(get_current_user)):
    """Every cycle that has no clip, with its full details and the reason."""
    _ensure_tables()
    d0, d1 = _range(date_from, date_to)
    with get_conn() as conn:
        cur = dict_cursor(conn)
        lines = _scope(cur, user, line_id)
        total, rows = _missing_rows(cur, d0, d1, shift, lines, machine_key, reason,
                                    page_size, (page - 1) * page_size)
    return {"total": total, "page": page, "page_size": page_size, "rows": rows}


@router.get("/cameras")
def cameras(user=Depends(get_current_user)):
    """Each camera's current state (latest tracker sample) and where it is bound."""
    _ensure_tables()
    with get_conn() as conn:
        cur = dict_cursor(conn)
        allowed = set(_scope(cur, user, None))
        lines, subs = _machines(cur)
        where = {}
        for ln in lines:
            if ln.get("cam") and ln["id"] in allowed:
                where.setdefault(ln["cam"], []).append(f"{ln['line_name']} · {ln.get('mname') or 'Final Inspection'}")
        lname = {ln["id"]: ln["line_name"] for ln in lines}
        for sp in subs.values():
            if sp.get("cam") and sp.get("line_id") in allowed and sp.get("line_id") in lname:
                where.setdefault(sp["cam"], []).append(f"{lname[sp['line_id']]} · {sp.get('mname') or 'Sub'}")
        cur.execute("""SELECT DISTINCT ON (camera_id) camera_id, ip, state, subnet_down, from_ts, to_ts
                         FROM mes_vcov_cam_state
                        WHERE to_ts > now() - interval '1 day'
                        ORDER BY camera_id, to_ts DESC""")
        rows = []
        for r in cur.fetchall():
            if r["camera_id"] not in where and user.get("role") != "admin":
                continue
            rows.append({"camera_id": r["camera_id"], "ip": r["ip"], "state": r["state"],
                         "subnet_down": r["subnet_down"],
                         "since": r["from_ts"].astimezone().isoformat(),
                         "last_seen": r["to_ts"].astimezone().isoformat(),
                         "bound_to": where.get(r["camera_id"], [])})
    order = {"cms_down": 0, "camera_offline": 1, "camera_hung": 2, "recording": 3}
    rows.sort(key=lambda x: (order.get(x["state"], 4), x["ip"] or ""))
    return {"cameras": rows}


@router.get("/agent")
def agent(user=Depends(get_current_user)):
    """Video Agent status and its findings (open + closed in the last 24 h)."""
    _ensure_tables()
    with get_conn() as conn:
        cur = dict_cursor(conn)
        cur.execute("SELECT * FROM mes_vcov_agent WHERE id = 1")
        st = cur.fetchone() or {}
        cur.execute("""SELECT f.*, l.line_name FROM mes_vcov_findings f
                         LEFT JOIN mes_lines l ON l.id = f.line_id
                        WHERE f.closed_at IS NULL OR f.closed_at > now() - interval '1 day'
                        ORDER BY (f.closed_at IS NULL) DESC, f.opened_at DESC LIMIT 300""")
        fs = cur.fetchall()
    iso = lambda v: v.astimezone().isoformat() if v else None
    return {"status": {k: iso(st.get(k)) for k in
                       ("started_at", "tracking_since", "last_sample", "last_eval",
                        "last_agent", "last_notify")} | {"last_error": st.get("last_error"),
                                                          "tracker_running": _STARTED or None},
            "findings": [{"id": f["id"], "kind": f["kind"], "line_name": f["line_name"],
                          "camera_id": f["camera_id"], "message": f["message"],
                          "opened_at": iso(f["opened_at"]), "last_seen": iso(f["last_seen"]),
                          "closed_at": iso(f["closed_at"])} for f in fs],
            "rules": [
                "CMS_DOWN — the CMS does not answer for 2+ minutes",
                "CAMERA_DOWN — a camera whose machine produced in the last 15 min has been offline or hung for 10+ min",
                f"LOW_COVERAGE — a line's judged-hour coverage is below 80% (20+ cycles); clears at 90%",
                "Read-only: the agent never restarts, stops or changes anything; admins get one Inbox digest at most every 15 min"]}


@router.get("/export")
def export(date_from: Optional[str] = Query(None), date_to: Optional[str] = Query(None),
           shift: str = Query("ALL"), line_id: Optional[int] = Query(None),
           reason: Optional[str] = Query(None), user=Depends(get_current_user)):
    """Excel: line summary + every missing cycle (up to 50,000 rows)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    s = summary(date_from, date_to, shift, line_id, user)
    d0, d1 = _range(date_from, date_to)
    with get_conn() as conn:
        cur = dict_cursor(conn)
        lines = _scope(cur, user, line_id)
        _total, rows = _missing_rows(cur, d0, d1, shift, lines, None, reason, 50000, 0)
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    clean = lambda v: ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v
    wb = Workbook()
    hfill, hfont = PatternFill("solid", fgColor="1E3A8A"), Font(bold=True, color="FFFFFF")

    def header(ws, cols):
        ws.append(cols)
        for c in ws[ws.max_row]:
            c.fill, c.font = hfill, hfont

    ws = wb.active
    ws.title = "Lines"
    ws.append([f"Video Coverage  {s['date_from']} to {s['date_to']}  Shift: {s['shift']}"])
    ws.append([f"Cycles {s['kpis']['cycles']}  Clips {s['kpis']['clips']}  "
               f"Missing {s['kpis']['missing']}  Coverage {s['kpis']['coverage']}%"])
    ws.append([])
    rlabels = list(REASONS)
    header(ws, ["Zone", "Line", "Cycles", "With clip", "Missing", "Coverage %"] +
           [REASONS[k] for k in rlabels])
    for ln in s["lines"]:
        ws.append([ln["zone_name"], ln["line_name"], ln["cycles"], ln["clips"], ln["missing"],
                   ln["coverage"]] + [ln["reasons"].get(k, 0) for k in rlabels])
    ws2 = wb.create_sheet("Missing cycles")
    header(ws2, ["Date", "Shift", "Line", "Machine", "Cycle #", "Part code", "OK/NG",
                 "Start", "End", "CT (s)", "Camera", "Camera IP", "Reason", "Detail"])
    for r in rows:
        ws2.append([clean(v) for v in (
                    r["record_date"], r["shift"], r["line_name"], r["machine_name"], r["cycle_seq"],
                    r["part_code"], "NG" if r["is_ng"] else "OK",
                    (r["ts_start"] or "")[11:19], r["ts_end"][11:19], r["ct"],
                    r["camera_id"], r["camera_ip"], r["reason_label"], r["detail"])])
    for wsx, widths in ((ws, [14, 22, 9, 10, 9, 11] + [14] * len(rlabels)),
                        (ws2, [11, 7, 18, 30, 8, 24, 7, 10, 10, 8, 38, 15, 34, 48])):
        for i, wdt in enumerate(widths):
            wsx.column_dimensions[chr(65 + i) if i < 26 else "A"].width = wdt
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"video_coverage_{s['date_from']}_{s['date_to']}_{s['shift']}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})


# ── Camera Status API ───────────────────────────────────────────────────────
@router.get("/camera-status")
def camera_status(user=Depends(get_current_user)):
    """Zone → line → camera: online / hung / offline, clips cut or not, and why."""
    _ensure_tables()
    with get_conn() as conn:
        cur = dict_cursor(conn)
        allowed = set(_scope(cur, user, None))
        return _camera_status(cur, allowed, include_unassigned=user.get("role") == "admin")


def _cam_where_lines(cur, user, line_id):
    """{camera_id: 'Line · machine, …'} for the cameras the caller may see."""
    allowed = set(_scope(cur, user, line_id))
    lines, subs = _machines(cur)
    lname = {ln["id"]: ln["line_name"] for ln in lines}
    where = {}
    for ln in lines:
        if ln.get("cam") and ln["id"] in allowed:
            where.setdefault(ln["cam"], []).append(f"{ln['line_name']} · {ln.get('mname') or 'Final Inspection'}")
    for sp in subs.values():
        if sp.get("cam") and sp.get("line_id") in allowed and sp.get("line_id") in lname:
            where.setdefault(sp["cam"], []).append(f"{lname[sp['line_id']]} · {sp.get('mname') or 'Sub'}")
    return where


def _camera_log_rows(cur, user, d0, d1, line_id, state, limit, offset, q=None):
    where = _cam_where_lines(cur, user, line_id)
    cams = list(where)
    if not cams:
        return 0, []
    w = ["camera_id = ANY(%s)", "from_ts < %s", "to_ts >= %s"]
    a = [cams, datetime.combine(d1 + timedelta(days=1), datetime.min.time()).astimezone(),
         datetime.combine(d0, datetime.min.time()).astimezone()]
    if state:
        w.append("state = %s")
        a.append(state)
    if q and q.strip():
        w.append("(camera_id ILIKE %s OR ip ILIKE %s)")
        a += [f"%{q.strip()}%", f"%{q.strip()}%"]
    cur.execute(f"SELECT count(*) n FROM mes_vcov_cam_state WHERE {' AND '.join(w)}", a)
    total = int(cur.fetchone()["n"])
    cur.execute(f"""SELECT camera_id, ip, state, subnet_down, from_ts, to_ts
                      FROM mes_vcov_cam_state WHERE {' AND '.join(w)}
                     ORDER BY from_ts DESC LIMIT %s OFFSET %s""", a + [limit, offset])
    rows = []
    for r in cur.fetchall():
        rows.append({"camera_id": r["camera_id"], "ip": r["ip"], "state": r["state"],
                     "subnet_down": r["subnet_down"],
                     "from": r["from_ts"].astimezone().isoformat(),
                     "to": r["to_ts"].astimezone().isoformat(),
                     "minutes": round((r["to_ts"] - r["from_ts"]).total_seconds() / 60.0, 1),
                     "bound_to": where.get(r["camera_id"], [])})
    return total, rows


@router.get("/camera-log")
def camera_log(date_from: Optional[str] = Query(None), date_to: Optional[str] = Query(None),
               line_id: Optional[int] = Query(None), state: Optional[str] = Query(None),
               q: Optional[str] = Query(None, description="camera id or IP contains"),
               page: int = Query(1, ge=1), page_size: int = Query(200, ge=20, le=1000),
               user=Depends(get_current_user)):
    """Every camera state period (recording / hung / offline / CMS down) — the camera log."""
    _ensure_tables()
    d0, d1 = _range(date_from, date_to)
    with get_conn() as conn:
        cur = dict_cursor(conn)
        total, rows = _camera_log_rows(cur, user, d0, d1, line_id, state,
                                       page_size, (page - 1) * page_size, q)
    return {"total": total, "page": page, "page_size": page_size, "rows": rows}


def _line_log_rows(cur, user, d0, d1, line_id, limit):
    lines = _scope(cur, user, line_id)
    cur.execute("""SELECT * FROM mes_vcov_line_snap
                    WHERE line_id = ANY(%s) AND ts >= %s AND ts < %s
                    ORDER BY ts DESC, zone_name, line_name LIMIT %s""",
                (lines, datetime.combine(d0, datetime.min.time()).astimezone(),
                 datetime.combine(d1 + timedelta(days=1), datetime.min.time()).astimezone(), limit))
    return [{**{k: r[k] for k in ("line_id", "zone_name", "line_name", "cameras", "online", "hung",
                                  "offline", "cms_down", "cutting", "not_cutting", "idle")},
             "ts": r["ts"].astimezone().isoformat()} for r in cur.fetchall()]


@router.get("/line-log")
def line_log(date_from: Optional[str] = Query(None), date_to: Optional[str] = Query(None),
             line_id: Optional[int] = Query(None), limit: int = Query(2000, ge=10, le=20000),
             user=Depends(get_current_user)):
    """Each line's camera counts every 5 minutes — the line log."""
    _ensure_tables()
    _ensure_status_tables()
    d0, d1 = _range(date_from, date_to)
    with get_conn() as conn:
        cur = dict_cursor(conn)
        return {"rows": _line_log_rows(cur, user, d0, d1, line_id, limit)}


@router.get("/camera-log/export")
def camera_log_export(date_from: Optional[str] = Query(None), date_to: Optional[str] = Query(None),
                      line_id: Optional[int] = Query(None), state: Optional[str] = Query(None),
                      user=Depends(get_current_user)):
    """Excel: live camera status + camera log + line log for the range."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    _ensure_tables()
    _ensure_status_tables()
    d0, d1 = _range(date_from, date_to)
    with get_conn() as conn:
        cur = dict_cursor(conn)
        allowed = set(_scope(cur, user, line_id))
        live = _camera_status(cur, allowed)
        _t, clog = _camera_log_rows(cur, user, d0, d1, line_id, state, 50000, 0)
        llog = _line_log_rows(cur, user, d0, d1, line_id, 20000)
    clean = lambda v: ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v
    lbl = {"recording": "Recording", "camera_hung": "Hung (ping OK, no video)",
           "camera_offline": "Offline (no ping)", "cms_down": "CMS down"}
    wb = Workbook()
    hfill, hfont = PatternFill("solid", fgColor="1E3A8A"), Font(bold=True, color="FFFFFF")

    def header(ws, cols):
        ws.append(cols)
        for c in ws[ws.max_row]:
            c.fill, c.font = hfill, hfont
    ws = wb.active
    ws.title = "Camera status (live)"
    header(ws, ["Zone", "Line", "Camera IP", "Machine(s)", "State", "Since", "Cycles (3-20 min ago)",
                "Clips", "Clip %", "Status", "Reason"])
    for z in live["zones"]:
        for L in z["lines"]:
            for c in L["cameras_list"]:
                ws.append([clean(v) for v in (z["zone_name"], L["line_name"], c["ip"], ", ".join(c["machines"]),
                           lbl.get(c["state"], c["state"] or "—"), (c["since"] or "")[:16].replace("T", " "),
                           c["cycles"], c["clips"], c["clip_pct"],
                           {"cutting": "Cutting", "not_cutting": "Not cutting", "idle": "Idle"}[c["status"]],
                           c["reason"] or "")])
    ws2 = wb.create_sheet("Camera log")
    header(ws2, ["From", "To", "Minutes", "State", "Camera IP", "Camera", "Used by"])
    for r in clog:
        ws2.append([clean(v) for v in (r["from"][:19].replace("T", " "), r["to"][:19].replace("T", " "), r["minutes"],
                    lbl.get(r["state"], r["state"]) + (" — whole subnet down" if r["subnet_down"] else ""),
                    r["ip"], r["camera_id"], ", ".join(r["bound_to"]))])
    ws3 = wb.create_sheet("Line log (5 min)")
    header(ws3, ["Time", "Zone", "Line", "Cameras", "Online", "Hung", "Offline", "CMS down",
                 "Cutting", "Not cutting", "Idle"])
    for r in llog:
        ws3.append([r["ts"][:16].replace("T", " "), r["zone_name"], r["line_name"], r["cameras"], r["online"],
                    r["hung"], r["offline"], r["cms_down"], r["cutting"], r["not_cutting"], r["idle"]])
    for wsx, widths in ((ws, [14, 18, 15, 40, 24, 17, 12, 8, 8, 12, 34]),
                        (ws2, [20, 20, 9, 34, 15, 44, 60]),
                        (ws3, [17, 14, 20, 9, 8, 7, 8, 9, 8, 11, 6])):
        for i, wdt in enumerate(widths):
            wsx.column_dimensions[chr(65 + i)].width = wdt
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"camera_status_{d0.isoformat()}_{d1.isoformat()}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})
