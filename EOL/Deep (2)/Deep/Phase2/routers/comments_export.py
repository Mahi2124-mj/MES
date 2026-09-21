"""routers/comments_export.py — Comments History as a real Excel workbook.

POST /api/comments-export/xlsx

Operator: "export karta hoon to excel format me chahiye" — the page used to save
an MHTML web-archive (.mht).  It carried the charts, but it is not a
spreadsheet: it opened in a browser, or in Excel only after a format warning,
and nothing in it could be sorted or filtered as a date.

The page sends exactly what it is showing — the rows after every filter,
including the client-side zone and cycle-time filters the API never sees — and
this only lays them out.  Nothing is re-queried, so the file cannot disagree
with the screen it was exported from.

    kind = "full"    Comments + Pareto + Summary sheets, both charts as pictures
    kind = "pareto"  the Pareto sheet on its own (the Pareto tab's own button)

Dates and times are written as real Excel values with a display format, so the
file shows the same "19-09-2026 / 11:13:27 AM" the page shows and still sorts
and filters as dates.
"""

import base64
import io
import re
from datetime import datetime
from typing import Any, List, Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel

from auth import get_current_user

router = APIRouter(prefix="/api/comments-export", tags=["comments-export"])

MAX_ROWS = 200_000               # a year of comments is ~40k; this is a guard, not a quota
MAX_PNG_BYTES = 6 * 1024 * 1024

HEADER_FILL = "DBE4F0"
DATE_FMT = "DD-MM-YYYY"
TIME_FMT = "hh:mm:ss AM/PM"


class Col(BaseModel):
    key: str
    header: str
    type: str = "text"           # text | date | time | number
    width: Optional[float] = None


class Chart(BaseModel):
    title: str
    png: str                     # data:image/png;base64,...


class ParetoItem(BaseModel):
    name: str
    count: float = 0
    pct: Optional[float] = None
    cum_pct: Optional[float] = None


class ExportReq(BaseModel):
    kind: str = "full"
    title: str = "Comments History"
    filters: str = ""
    filename: str = "comments-history.xlsx"
    summary: List[List[Any]] = []
    columns: List[Col] = []
    rows: List[List[Any]] = []
    pareto: List[ParetoItem] = []
    pareto_scope: str = ""
    charts: List[Chart] = []


def _safe_name(name: str) -> str:
    base = re.sub(r"[^A-Za-z0-9._-]+", "_", name or "export").strip("._") or "export"
    return base if base.lower().endswith(".xlsx") else base + ".xlsx"


def _as_date(v):
    try:
        return datetime.strptime(str(v)[:10], "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None


def _as_time(v):
    try:
        return datetime.strptime(str(v)[:8], "%H:%M:%S").time()
    except (TypeError, ValueError):
        return None


def _as_number(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _png(data_url: str):
    """A chart picture for openpyxl, or None if the data is unusable."""
    from openpyxl.drawing.image import Image as XLImage
    try:
        raw = base64.b64decode(str(data_url).split(",", 1)[-1], validate=False)
    except Exception:
        return None
    if not raw or len(raw) > MAX_PNG_BYTES or not raw.startswith(b"\x89PNG"):
        return None
    try:
        img = XLImage(io.BytesIO(raw))
    except Exception:
        return None
    # The page draws charts at 1000 x 420; 860 wide fits a laptop-width sheet.
    if img.width:
        scale = 860 / img.width
        img.width, img.height = int(img.width * scale), int(img.height * scale)
    return img


def _styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    thin = Side(style="thin", color="B8C4D6")
    return {
        "title":  Font(bold=True, size=14, color="0F172A"),
        "sub":    Font(italic=True, size=10, color="475569"),
        "hfont":  Font(bold=True, color="0F172A"),
        "hfill":  PatternFill("solid", fgColor=HEADER_FILL),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "wrap":   Alignment(wrap_text=True, vertical="top"),
        "top":    Alignment(vertical="top"),
        "hcell":  Alignment(wrap_text=True, vertical="center"),
    }


def _heading(ws, st, title, filters):
    ws["A1"] = title
    ws["A1"].font = st["title"]
    if filters:
        ws["A2"] = filters
        ws["A2"].font = st["sub"]
    ws["A3"] = "Exported " + datetime.now().strftime("%d-%m-%Y %I:%M %p")
    ws["A3"].font = st["sub"]


def _header_row(ws, st, row, headers):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill, c.border, c.alignment = st["hfont"], st["hfill"], st["border"], st["hcell"]


def _comments_sheet(wb, st, req: ExportReq):
    from openpyxl.utils import get_column_letter
    ws = wb.create_sheet("Comments")
    _heading(ws, st, req.title, req.filters)
    cols = req.columns
    hdr = 5
    _header_row(ws, st, hdr, [c.header for c in cols])
    for r_i, row in enumerate(req.rows, hdr + 1):
        for c_i, col in enumerate(cols, 1):
            v = row[c_i - 1] if c_i - 1 < len(row) else None
            if col.type == "date":
                val, fmt = _as_date(v), DATE_FMT
            elif col.type == "time":
                val, fmt = _as_time(v), TIME_FMT
            elif col.type == "number":
                val, fmt = _as_number(v), "0.00"
            else:
                val, fmt = ("" if v is None else str(v)), None
            cell = ws.cell(row=r_i, column=c_i, value=val)
            cell.border = st["border"]
            cell.alignment = st["wrap"] if col.type == "text" and (col.width or 0) >= 30 else st["top"]
            if fmt:
                cell.number_format = fmt
    for c_i, col in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(c_i)].width = col.width or 14
    ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
    if cols:
        last = hdr + max(len(req.rows), 1)
        ws.auto_filter.ref = f"A{hdr}:{get_column_letter(len(cols))}{last}"
    return ws


def _pareto_sheet(wb, st, req: ExportReq, chart):
    ws = wb.create_sheet("Pareto")
    title = "Problem Pareto" + (f" — {req.pareto_scope}" if req.pareto_scope else "")
    _heading(ws, st, title, req.filters)
    hdr = 5
    _header_row(ws, st, hdr, ["Rank", "Problem", "Count", "Share %", "Cumulative %"])
    for i, it in enumerate(req.pareto, 1):
        r = hdr + i
        vals = [i, it.name, it.count, it.pct, it.cum_pct]
        fmts = ["0", None, "0", "0.0", "0.0"]
        for c_i, (v, f) in enumerate(zip(vals, fmts), 1):
            cell = ws.cell(row=r, column=c_i, value=v)
            cell.border = st["border"]
            cell.alignment = st["wrap"] if c_i == 2 else st["top"]
            if f:
                cell.number_format = f
    for letter, w in zip("ABCDE", (7, 46, 10, 10, 13)):
        ws.column_dimensions[letter].width = w
    ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
    if req.pareto:
        ws.auto_filter.ref = f"A{hdr}:E{hdr + len(req.pareto)}"
    # The chart sits beside the table so it is on screen as soon as the sheet
    # opens, however long the ranking runs (997 distinct problems in a week).
    if chart is not None:
        ws.add_image(chart, "G5")
    return ws


def _summary_sheet(wb, st, req: ExportReq, chart):
    ws = wb.create_sheet("Summary")
    _heading(ws, st, req.title + " — Summary", req.filters)
    hdr = 5
    _header_row(ws, st, hdr, ["Measure", "Value"])
    for i, pair in enumerate(req.summary, 1):
        k = pair[0] if len(pair) > 0 else ""
        v = pair[1] if len(pair) > 1 else ""
        for c_i, val in enumerate((k, v), 1):
            cell = ws.cell(row=hdr + i, column=c_i, value=val)
            cell.border = st["border"]
            if c_i == 2 and isinstance(val, (int, float)):
                cell.number_format = "#,##0"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 16
    if chart is not None:
        ws.add_image(chart, f"A{hdr + len(req.summary) + 3}")
    return ws


@router.post("/xlsx")
def comments_xlsx(req: ExportReq, user=Depends(get_current_user)):
    """Lay out exactly what the Comments History page is showing as .xlsx."""
    from openpyxl import Workbook

    if len(req.rows) > MAX_ROWS:
        raise HTTPException(413, f"{len(req.rows)} rows is more than one file should hold "
                                 f"— narrow the date range")
    if req.kind not in ("full", "pareto"):
        raise HTTPException(400, "kind must be 'full' or 'pareto'")

    charts = {c.title: c.png for c in req.charts}
    ct_chart = next((v for k, v in charts.items() if "cycle" in k.lower()), None)
    pa_chart = next((v for k, v in charts.items() if "pareto" in k.lower()), None)

    st = _styles()
    wb = Workbook()
    wb.remove(wb.active)
    if req.kind == "full":
        _comments_sheet(wb, st, req)
        _pareto_sheet(wb, st, req, _png(pa_chart) if pa_chart else None)
        _summary_sheet(wb, st, req, _png(ct_chart) if ct_chart else None)
    else:
        _pareto_sheet(wb, st, req, _png(pa_chart) if pa_chart else None)

    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{_safe_name(req.filename)}"'},
    )
